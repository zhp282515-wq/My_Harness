"""解析 XLSX 工作簿，并将每个非空工作表转换为 HTML 表格。"""

from datetime import date, datetime, time
from decimal import Decimal
from html import escape
from pathlib import Path
from typing import Any

from utils.document_models import XLSXDocument, XLSXSheet


def _cell_text(value: object) -> str:
    """将单元格值转成可读文本，保留日期和换行等信息。"""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (datetime, date, time)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, Decimal):
        return format(value, "f")
    return str(value)


def _worksheet_to_html(worksheet: Any) -> XLSXSheet | None:
    from openpyxl.utils import get_column_letter

    # openpyxl 的 Worksheet 接口由对象提供；保持辅助函数的公开签名不暴露该类型。
    ws = worksheet
    populated = [
        (cell.row, cell.column)
        for row in ws.iter_rows()
        for cell in row
        if cell.value is not None
    ]
    if not populated:
        return None

    bounds = populated.copy()
    for merged_range in ws.merged_cells.ranges:
        bounds.append((merged_range.min_row, merged_range.min_col))
        bounds.append((merged_range.max_row, merged_range.max_col))

    min_row = min(row for row, _ in bounds)
    max_row = max(row for row, _ in bounds)
    min_col = min(column for _, column in bounds)
    max_col = max(column for _, column in bounds)

    # 记录每行的合并区间。合并范围内仅左上角锚点输出单元格，其余位置由 rowspan/colspan 覆盖。
    merges_by_row: dict[int, list[tuple[int, int, int, int, int, int]]] = {}
    for merged_range in ws.merged_cells.ranges:
        span = (
            merged_range.min_col,
            merged_range.max_col,
            merged_range.min_row,
            merged_range.min_col,
            merged_range.max_row - merged_range.min_row + 1,
            merged_range.max_col - merged_range.min_col + 1,
        )
        for row_index in range(merged_range.min_row, merged_range.max_row + 1):
            merges_by_row.setdefault(row_index, []).append(span)
    for row_merges in merges_by_row.values():
        row_merges.sort(key=lambda item: item[0])

    cell_range = (
        f"{get_column_letter(min_col)}{min_row}:"
        f"{get_column_letter(max_col)}{max_row}"
    )
    html_parts = ["<table>", f"<caption>{escape(ws.title, quote=False)}</caption>", "<tbody>"]

    for row_index in range(min_row, max_row + 1):
        html_parts.append("<tr>")
        row_merges = merges_by_row.get(row_index, [])
        merge_index = 0
        column = min_col

        while column <= max_col:
            while merge_index < len(row_merges) and row_merges[merge_index][1] < column:
                merge_index += 1

            current_merge = (
                row_merges[merge_index]
                if merge_index < len(row_merges)
                and row_merges[merge_index][0] <= column <= row_merges[merge_index][1]
                else None
            )

            if current_merge is not None:
                start_col, end_col, anchor_row, anchor_col, rowspan, colspan = current_merge
                if row_index != anchor_row or column != anchor_col:
                    column = end_col + 1
                    continue
                attrs = []
                if rowspan > 1:
                    attrs.append(f'rowspan="{rowspan}"')
                if colspan > 1:
                    attrs.append(f'colspan="{colspan}"')
                attr_text = f" {' '.join(attrs)}" if attrs else ""
                value = ws.cell(row=anchor_row, column=anchor_col).value
                cell_ref = f"{get_column_letter(anchor_col)}{anchor_row}"
                html_parts.append(
                    f'<td data-cell="{cell_ref}"{attr_text}>'
                    f"{escape(_cell_text(value), quote=False)}</td>"
                )
                column = end_col + 1
                continue

            value = ws.cell(row=row_index, column=column).value
            cell_ref = f"{get_column_letter(column)}{row_index}"
            html_parts.append(
                f'<td data-cell="{cell_ref}">'
                f"{escape(_cell_text(value), quote=False)}</td>"
            )
            column += 1

        html_parts.append("</tr>")

    html_parts.extend(["</tbody>", "</table>"])
    return XLSXSheet(
        html="".join(html_parts),
        name=ws.title,
        cell_range=cell_range,
    )


def parse_xlsx(
    file_path: str | Path,
    *,
    data_only: bool = False,
) -> XLSXDocument:
    """Parse one workbook into a document containing its non-empty worksheets.

    Parsing preserves worksheet order, cell positions, and merged-cell spans without
    splitting tables into retrieval chunks. By default formula cells are represented
    by their formula text; pass ``data_only=True`` to read cached formula results.
    Empty worksheets are skipped. The workbook is always closed before returning.
    """
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ImportError("解析 XLSX 需要安装 openpyxl：pip install openpyxl") from exc

    path = Path(file_path)
    workbook = load_workbook(path, data_only=data_only, read_only=False, keep_links=False)
    sheets: list[XLSXSheet] = []

    try:
        for worksheet in workbook.worksheets:
            sheet = _worksheet_to_html(worksheet)
            if sheet is not None:
                sheets.append(sheet)
    finally:
        workbook.close()

    return XLSXDocument(
        source_file=str(path.resolve()),
        sheets=tuple(sheets),
    )


__all__ = ["XLSXDocument", "XLSXSheet", "parse_xlsx"]
