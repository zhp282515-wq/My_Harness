"""将 DOCX 按页解析为保持原文顺序的 HTML 内容。"""

from __future__ import annotations

from html import escape
import os
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
from dataclasses import dataclass
from uuid import uuid4

from utils.document_models import (
    DOCXDocument,
    DOCXPage,
)
from utils.html_utils import local_file_uri
from utils.logger_tool import logger


_WRAPPERS = {"sdt", "sdtContent", "customXml", "ins", "del", "moveFrom", "moveTo", "smartTag"}
_IMAGE_MIME_EXTENSIONS = {
    "image/jpeg": ".jpg",
    "image/png": ".png",
    "image/gif": ".gif",
    "image/bmp": ".bmp",
    "image/tiff": ".tiff",
    "image/svg+xml": ".svg",
}


@dataclass(frozen=True, slots=True)
class _ImageAsset:
    path: str
    description: str
    alt_text: str = ""


def _local_name(element) -> str:
    return element.tag.rsplit("}", 1)[-1]


def _inline_text(value: str) -> str:
    """移除文字字段中的换行符，避免 content 含原始换行。"""
    return re.sub(r"[\r\n]+", " ", value)


def _iter_effective_nodes(element):
    """展开 OOXML，并对 AlternateContent 只采用一个兼容分支。"""
    if _local_name(element) == "AlternateContent":
        choice = next(
            (child for child in element if _local_name(child) == "Choice"),
            None,
        )
        fallback = next(
            (child for child in element if _local_name(child) == "Fallback"),
            None,
        )
        branch = choice if choice is not None else fallback
        if branch is not None:
            yield from _iter_effective_nodes(branch)
        return

    yield element
    for child in element:
        yield from _iter_effective_nodes(child)


def _iter_blocks(parent):
    """保持 OOXML 中段落和表格的正文顺序，递归展开内容控件包装。"""
    for child in parent.iterchildren():
        name = _local_name(child)
        if name in {"p", "tbl"}:
            yield child
        elif name in _WRAPPERS:
            yield from _iter_blocks(child)


def _iter_paragraph_text(element) -> str:
    parts: list[str] = []
    for node in _iter_effective_nodes(element):
        name = _local_name(node)
        if name == "t" and node.text:
            parts.append(node.text)
        elif name in {"drawing", "pict"} and _image_relationships(node):
            parts.append("[图片]")
        elif name == "tab":
            parts.append("\t")
        elif name in {"br", "cr"}:
            parts.append("\n")
    return "".join(parts)


def _image_relationships(element) -> list[str]:
    rel_ids: list[str] = []
    seen: set[str] = set()
    for node in _iter_effective_nodes(element):
        name = _local_name(node)
        if name in {"blip", "imagedata"}:
            rel_id = node.get(
                "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
            ) or node.get(
                "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
            )
            if rel_id and rel_id not in seen:
                rel_ids.append(rel_id)
                seen.add(rel_id)
    return rel_ids


def _has_visual_content(element) -> bool:
    return any(
        _local_name(node) in {"drawing", "pict", "object"}
        for node in _iter_effective_nodes(element)
    )


def _alt_text(element, rel_id: str) -> str:
    for drawing in _iter_effective_nodes(element):
        if _local_name(drawing) not in {"drawing", "pict"}:
            continue
        if rel_id not in _image_relationships(drawing):
            continue
        for node in drawing.iter():
            if _local_name(node) in {"docPr", "cNvPr"}:
                return node.get("descr") or node.get("title") or ""
    return ""


def _describe_image(image_path: Path) -> str:
    from utils.vlm import parse_image

    return parse_image(str(image_path))


def _save_image_part(image_part, output_dir: Path, sequence: int) -> Path:
    """提取 DOCX 图片；常见栅格格式统一转成约定的 PNG 文件。"""
    import pymupdf

    try:
        pixmap = pymupdf.Pixmap(image_part.blob)
        if pixmap.colorspace is None or pixmap.colorspace.n > 3:
            pixmap = pymupdf.Pixmap(pymupdf.csRGB, pixmap)
        output = output_dir / f"image_{sequence:04d}.png"
        pixmap.save(output)
        return output
    except Exception:
        # 少见矢量格式（例如 EMF）无法由 MuPDF 解码时保留其真实格式，
        # 不把原始二进制误存成 PNG。
        mime_type = image_part.content_type or "application/octet-stream"
        extension = (
            _IMAGE_MIME_EXTENSIONS.get(mime_type)
            or Path(image_part.partname).suffix
            or ".bin"
        )
        output = output_dir / f"image_{sequence:04d}{extension}"
        output.write_bytes(image_part.blob)
        logger.warning(
            "[DOCX] 图片无法转换为 PNG，保留原始格式：%s (%s)",
            output,
            mime_type,
        )
        return output


def _export_docx_pdf(path: Path, output_dir: Path) -> Path:
    """按本机可用的 Writer 自动化工具将 DOCX 排版并导出 PDF。"""
    if os.name == "nt":
        app = None
        document = None
        try:
            import win32com.client

            pdf_path = output_dir / f"{path.stem}.pdf"
            try:
                app = win32com.client.DispatchEx("Word.Application")
            except Exception:
                app = win32com.client.DispatchEx("KWPS.Application")
            app.Visible = False
            app.DisplayAlerts = 0
            document = app.Documents.Open(str(path.resolve()), ReadOnly=True, AddToRecentFiles=False)
            document.ExportAsFixedFormat(str(pdf_path), 17)
            return pdf_path
        except ImportError:
            logger.info("[DOCX] pywin32 不可用，尝试 LibreOffice 导出分页 PDF")
        except Exception:
            logger.exception("[DOCX] Word/WPS 导出 PDF 失败，尝试 LibreOffice")
        finally:
            if document is not None:
                try:
                    document.Close(False)
                except Exception:
                    pass
            if app is not None:
                try:
                    app.Quit()
                except Exception:
                    pass

    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if executable is None:
        raise RuntimeError(
            "DOCX 按页解析需要 Word/WPS 自动化或 LibreOffice；"
            "Windows 请安装 pywin32 并确保 Word/WPS 可用，其他系统请安装 LibreOffice"
        )
    result = subprocess.run(
        [executable, "--headless", "--convert-to", "pdf", "--outdir", str(output_dir), str(path)],
        capture_output=True,
        text=True,
        timeout=180,
        check=False,
    )
    pdf_path = output_dir / f"{path.stem}.pdf"
    if result.returncode != 0 or not pdf_path.is_file():
        raise RuntimeError(f"LibreOffice 导出 DOCX PDF 失败：{result.stderr[-1000:]}")
    return pdf_path


def _page_boundaries(path: Path, pdf_path: Path):
    """按 DOCX 原始顺序将段落、表格行和表格内段落锚定到 PDF 页码。"""
    import pymupdf as fitz
    from docx import Document as OpenDocument

    doc = OpenDocument(path)
    text_occurrences: list[tuple[str, int, int | None]] = []
    for sequence, element in enumerate(_iter_blocks(doc.element.body), start=1):
        if _local_name(element) == "p":
            value = _iter_paragraph_text(element).strip()
            if value:
                text_occurrences.append((value, sequence, None))
        else:
            for row_index, row in enumerate(element.tr_lst):
                for paragraph in (node for node in row.iter() if _local_name(node) == "p"):
                    value = _iter_paragraph_text(paragraph).strip()
                    if value:
                        text_occurrences.append((value, sequence, row_index))

    paragraph_page: dict[int, int] = {}
    table_row_pages: dict[tuple[int, int], int] = {}
    with fitz.open(pdf_path) as pdf:
        page_texts = [page.get_text("text") for page in pdf]
        normalized_pages = [re.sub(r"\s+", "", text) for text in page_texts]
        cursor_page = 0
        cursor_by_page = [0] * len(page_texts)
        for value, sequence, row_index in text_occurrences:
            normalized = re.sub(r"\s+", "", value)
            if len(normalized) < 3:
                continue
            # 限长锚点可规避超长段落行断裂；从当前位置向后保持文档顺序搜索。
            anchor = normalized[: min(48, len(normalized))]
            match_page = None
            match_pos = None
            for page_index in range(cursor_page, len(page_texts)):
                found = normalized_pages[page_index].find(
                    anchor, cursor_by_page[page_index] if page_index == cursor_page else 0
                )
                if found >= 0:
                    match_page, match_pos = page_index, found
                    break
            if match_page is not None:
                if row_index is None:
                    paragraph_page[sequence] = match_page + 1
                else:
                    table_row_pages.setdefault((sequence, row_index), match_page + 1)
                cursor_page = match_page
                cursor_by_page[match_page] = int(match_pos or 0) + len(anchor)
        return paragraph_page, table_row_pages, len(pdf)


def _word_page_boundaries(path: Path):
    """用 Word/WPS 的排版引擎读取段落页码和每个表格行的页码。"""
    if os.name != "nt":
        return None
    try:
        import win32com.client
    except ImportError:
        return None
    from docx import Document as OpenDocument

    word = document = None
    try:
        try:
            word = win32com.client.DispatchEx("Word.Application")
        except Exception:
            word = win32com.client.DispatchEx("KWPS.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        document = word.Documents.Open(
            str(path.resolve()), ReadOnly=True, AddToRecentFiles=False
        )
        page_count = max(int(document.ComputeStatistics(2)), 1)  # wdStatisticPages
        paragraph_pages: dict[int, int] = {}
        table_row_pages: dict[tuple[int, int], int] = {}
        cursor = int(document.Content.Start)
        com_table_cursor = 1
        docx = OpenDocument(path)
        for source_sequence, element in enumerate(_iter_blocks(docx.element.body), start=1):
            if _local_name(element) == "p":
                anchors = [
                    node.text.strip()
                    for node in element.iter()
                    if _local_name(node) == "t"
                    and node.text
                    and len(re.sub(r"\s+", "", node.text)) >= 3
                ]
                if not anchors:
                    if _image_relationships(element):
                        try:
                            paragraph_pages[source_sequence] = int(
                                document.Range(
                                    cursor, min(cursor + 1, document.Content.End)
                                ).Information(3)
                            )
                        except Exception:
                            pass
                    continue
                found_range = None
                anchor = max(anchors, key=lambda item: len(re.sub(r"\s+", "", item)))
                for candidate in (anchor[:60], anchor[:32], anchor[:12]):
                    search_range = document.Range(cursor, document.Content.End)
                    finder = search_range.Find
                    finder.ClearFormatting()
                    finder.Text = candidate
                    finder.Forward = True
                    finder.Wrap = 0  # wdFindStop
                    finder.MatchWildcards = False
                    finder.MatchCase = False
                    try:
                        if finder.Execute():
                            found_range = search_range
                            break
                    except Exception:
                        continue
                if found_range is not None:
                    try:
                        paragraph_pages[source_sequence] = int(found_range.Information(3))
                        cursor = max(cursor, int(found_range.End))
                    except Exception:
                        logger.debug("[DOCX] 段落页码读取失败，块序号 %d", source_sequence)
            else:
                expected = "".join(
                    re.sub(r"\s+", "", node.text)
                    for node in element.iter()
                    if _local_name(node) == "t" and node.text
                )
                match_table = None
                for index in range(com_table_cursor, document.Tables.Count + 1):
                    candidate_table = document.Tables.Item(index)
                    try:
                        candidate = re.sub(r"[\s\x07]+", "", candidate_table.Range.Text)
                    except Exception:
                        continue
                    if not expected or expected[: min(24, len(expected))] in candidate or candidate[: min(24, len(candidate))] in expected:
                        if int(candidate_table.Range.Start) >= cursor:
                            match_table = candidate_table
                            com_table_cursor = index + 1
                            break
                if match_table is not None:
                    row_search_start = int(match_table.Range.Start)
                    for row_index, row in enumerate(element.tr_lst):
                        page_candidates: list[int] = []
                        for paragraph in (node for node in row.iter() if _local_name(node) == "p"):
                            anchors = [
                                node.text.strip()
                                for node in paragraph.iter()
                                if _local_name(node) == "t" and node.text
                                and len(re.sub(r"\s+", "", node.text)) >= 3
                            ]
                            if not anchors:
                                continue
                            anchor = max(anchors, key=lambda item: len(re.sub(r"\s+", "", item)))
                            paragraph_page = None
                            for candidate in (anchor[:48], anchor[:24], anchor[:10]):
                                finder_range = document.Range(
                                    row_search_start, int(match_table.Range.End)
                                )
                                finder = finder_range.Find
                                finder.ClearFormatting()
                                finder.Text = candidate
                                finder.Forward = True
                                finder.Wrap = 0
                                finder.MatchWildcards = False
                                finder.MatchCase = False
                                try:
                                    if finder.Execute():
                                        paragraph_page = int(finder_range.Information(3))
                                        row_search_start = max(
                                            row_search_start, int(finder_range.End)
                                        )
                                        break
                                except Exception:
                                    continue
                            if paragraph_page is not None:
                                page_candidates.append(paragraph_page)
                        row_page = (
                            min(page_candidates)
                            if page_candidates
                            else max(1, int(match_table.Range.Information(3)))
                        )
                        table_row_pages[(source_sequence, row_index)] = row_page
                    cursor = max(cursor, int(match_table.Range.End))
        return paragraph_pages, table_row_pages, page_count
    except Exception:
        logger.exception("[DOCX] Word/WPS 页码读取失败，改用 PDF 页面映射")
        return None
    finally:
        if document is not None:
            try:
                document.Close(False)
            except Exception:
                pass
        if word is not None:
            try:
                word.Quit()
            except Exception:
                pass


def _render_paragraph(paragraph, image_refs: dict[str, _ImageAsset]) -> str:
    """按 Word run/XML 顺序输出文字、换行、超链接、格式及图片描述。"""
    from docx.text.run import Run

    def render(element) -> str:
        name = _local_name(element)
        if name in {"pPr", "rPr", "sectPr"}:
            return ""
        if name == "AlternateContent":
            branch = next(
                (child for child in element if _local_name(child) == "Choice"),
                None,
            )
            if branch is None:
                branch = next(
                    (child for child in element if _local_name(child) == "Fallback"),
                    None,
                )
            return render(branch) if branch is not None else ""
        if name == "t":
            return escape(_inline_text(element.text or ""), quote=False)
        if name == "tab":
            return "&#9;"
        if name in {"br", "cr"}:
            return "<br>"
        if name in {"drawing", "pict"}:
            content = []
            for rel_id in _image_relationships(element):
                image = image_refs.get(rel_id)
                if image is None:
                    content.append("[图片资源不可用]")
                    continue
                content.append(
                    f'<img src="{escape(local_file_uri(image.path), quote=True)}" '
                    f'alt="{escape(_inline_text(image.alt_text or "图片"), quote=True)}" '
                    'style="max-width:100%;height:auto">'
                    f'<span class="image-description">'
                    f'{escape(_inline_text(image.description), quote=False)}</span>'
                )
            textbox = "".join(
                escape(_inline_text(node.text or ""), quote=False)
                for node in element.iter()
                if _local_name(node) == "t"
            )
            if textbox:
                content.append(f'<span class="docx-textbox">{textbox}</span>')
            return " ".join(content) or "[图形对象]"
        if name == "r":
            content = "".join(render(child) for child in element)
            run = Run(element, paragraph)
            if run.bold:
                content = f"<strong>{content}</strong>"
            if run.italic:
                content = f"<em>{content}</em>"
            if run.underline:
                content = f"<u>{content}</u>"
            if run.font.strike:
                content = f"<del>{content}</del>"
            if run.font.superscript:
                content = f"<sup>{content}</sup>"
            elif run.font.subscript:
                content = f"<sub>{content}</sub>"
            return content
        if name == "hyperlink":
            content = "".join(render(child) for child in element)
            rel_id = element.get(
                "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
            )
            relation = paragraph.part.rels.get(rel_id) if rel_id else None
            if relation is None:
                return content
            return f'<a href="{escape(relation.target_ref, quote=True)}">{content}</a>'
        return "".join(render(child) for child in element)

    content = "".join(render(child) for child in paragraph._p)
    styles = ["white-space: pre-wrap"]
    alignment = {0: "left", 1: "center", 2: "right", 3: "justify"}.get(paragraph.alignment)
    if alignment:
        styles.append(f"text-align: {alignment}")
    return f'<p style="{"; ".join(styles)}">{content}</p>' if content else ""


def _paragraph_list(paragraph):
    """识别表格单元格中的 Word 编号段落。"""
    properties = paragraph._p.pPr
    numbering = properties.numPr if properties is not None else None
    if numbering is None or numbering.numId is None:
        style_properties = paragraph.style._element.pPr
        numbering = style_properties.numPr if style_properties is not None else None
    if numbering is None or numbering.numId is None:
        return None

    level = int(numbering.ilvl.val) if numbering.ilvl is not None else 0
    num_id = int(numbering.numId.val)
    from docx.oxml.ns import qn

    numbering_part = paragraph.part.numbering_part.element
    num = next(
        (item for item in numbering_part.num_lst if int(item.get(qn("w:numId"))) == num_id),
        None,
    )
    number_format = "bullet"
    if num is not None:
        abstract_id = num.abstractNumId.val
        abstract = next(
            (item for item in numbering_part if _local_name(item) == "abstractNum"
             and item.get(qn("w:abstractNumId")) == str(abstract_id)),
            None,
        )
        if abstract is not None:
            level_xml = next(
                (item for item in abstract if _local_name(item) == "lvl"
                 and item.get(qn("w:ilvl")) == str(level)),
                None,
            )
            if level_xml is not None:
                format_xml = next(
                    (item for item in level_xml if _local_name(item) == "numFmt"),
                    None,
                )
                if format_xml is not None:
                    number_format = format_xml.get(qn("w:val"), "bullet")
    return ("unordered" if number_format == "bullet" else "ordered"), level


def _cell_content(cell, part, assets_dir: Path, image_start: int,
                  image_cache: dict[str, _ImageAsset]) -> tuple[str, list[_ImageAsset], int]:
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    pieces: list[str] = []
    images: list[_ImageAsset] = []
    next_image = image_start
    for element in _iter_blocks(cell._tc):
        if _local_name(element) == "p":
            paragraph = Paragraph(element, cell)
            local_images: list[_ImageAsset] = []
            for rel_id in _image_relationships(element):
                image = image_cache.get(rel_id)
                if image is None:
                    image_part = part.related_parts.get(rel_id)
                    if image_part is None:
                        continue
                    output = _save_image_part(
                        image_part, assets_dir, next_image
                    )
                    alt = _alt_text(element, rel_id)
                    try:
                        description = _describe_image(output)
                    except Exception:
                        logger.exception("[DOCX] 图片描述失败：%s", output)
                        description = alt or "[图片描述生成失败]"
                    image = _ImageAsset(str(output.resolve()), description, alt)
                    image_cache[rel_id] = image
                    next_image += 1
                local_images.append(image)
                if image not in images:
                    images.append(image)
            image_refs = {
                rel_id: image
                for rel_id, image in image_cache.items()
                if image in local_images
            }
            paragraph_html = _render_paragraph(paragraph, image_refs)
            list_info = _paragraph_list(paragraph)
            if list_info is not None:
                list_type, list_level = list_info
                tag = "ol" if list_type == "ordered" else "ul"
                paragraph_html = f'<{tag} data-list-level="{list_level}"><li>{paragraph_html}</li></{tag}>'
            pieces.append(paragraph_html)
        else:
            nested = Table(element, cell)
            html, nested_images, next_image = _table_html(
                nested, part, assets_dir, next_image, image_cache
            )
            pieces.append(html)
            images.extend(item for item in nested_images if item not in images)
    return "".join(pieces), images, next_image


def _table_html(table, part, assets_dir: Path, image_start: int,
                image_cache: dict[str, _ImageAsset],
                row_indices: list[int] | None = None) -> tuple[str, list[_ImageAsset], int]:
    from docx.table import _Cell

    rows: list[list[tuple[object, int, int, str | None]]] = []
    source_rows = table._tbl.tr_lst
    selected_indices = row_indices if row_indices is not None else list(range(len(source_rows)))
    for source_index in selected_indices:
        row = source_rows[source_index]
        properties = row.trPr
        before = int(properties.gridBefore.val) if properties is not None and properties.gridBefore is not None else 0
        items = []
        column = before
        for tc in row.tc_lst:
            props = tc.tcPr
            colspan = int(props.gridSpan.val) if props is not None and props.gridSpan is not None else 1
            merge = (props.vMerge.val or "continue") if props is not None and props.vMerge is not None else None
            items.append((tc, column, colspan, merge))
            column += colspan
        rows.append(items)

    result = ["<table><tbody>"]
    images: list[_ImageAsset] = []
    next_image = image_start
    for row_index, row in enumerate(rows):
        result.append("<tr>")
        for tc, column, colspan, merge in row:
            if merge == "continue":
                # 若合并起始单元格已在本 HTML 片段内输出，rowspan 已覆盖此格。
                # 跨页片段缺少起始格时则显式输出续格，避免丢失其内容。
                started_in_fragment = any(
                    any(
                        prior_column == column
                        and prior_colspan == colspan
                        and prior_merge == "restart"
                        for _, prior_column, prior_colspan, prior_merge in prior_row
                    )
                    for prior_row in rows[:row_index]
                )
                if started_in_fragment:
                    continue
                cell = _Cell(tc, table)
                content, cell_images, next_image = _cell_content(
                    cell, part, assets_dir, next_image, image_cache
                )
                images.extend(item for item in cell_images if item not in images)
                attrs = [f'data-cell="R{selected_indices[row_index] + 1}C{column + 1}"', 'data-merge-continuation="true"']
                if colspan > 1:
                    attrs.append(f'colspan="{colspan}"')
                result.append(f"<td {' '.join(attrs)}>{content}</td>")
                continue
            rowspan = 1
            if merge == "restart":
                for following in rows[row_index + 1:]:
                    continuation = next((item for item in following if item[1] == column and item[2] == colspan), None)
                    if continuation is None or continuation[3] != "continue":
                        break
                    rowspan += 1
            cell = _Cell(tc, table)
            content, cell_images, next_image = _cell_content(
                cell, part, assets_dir, next_image, image_cache
            )
            images.extend(item for item in cell_images if item not in images)
            attrs = [f'data-cell="R{selected_indices[row_index] + 1}C{column + 1}"']
            if rowspan > 1:
                attrs.append(f'rowspan="{rowspan}"')
            if colspan > 1:
                attrs.append(f'colspan="{colspan}"')
            result.append(f"<td {' '.join(attrs)}>{content}</td>")
        result.append("</tr>")
    result.append("</tbody></table>")
    return "".join(result), images, next_image


def _document_pages(path: Path, assets_dir: Path,
                    paragraph_pages: dict[int, int],
                    table_row_pages: dict[tuple[int, int], int], page_count: int):
    from docx import Document as OpenDocument
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = OpenDocument(path)
    page_contents: list[list[str]] = [[] for _ in range(page_count)]
    image_cache: dict[str, _ImageAsset] = {}
    image_sequence = 1
    seen_page = 1

    for source_sequence, element in enumerate(_iter_blocks(document.element.body), start=1):
        name = _local_name(element)
        if name == "p":
            paragraph = Paragraph(element, document)
            text = paragraph.text
            page_number = paragraph_pages.get(source_sequence, seen_page)
            page_number = min(max(page_number, 1), page_count)
            seen_page = max(seen_page, page_number)
            rel_ids = _image_relationships(element)
            has_visual_content = _has_visual_content(element)
            block_images: list[_ImageAsset] = []
            for rel_id in rel_ids:
                image = image_cache.get(rel_id)
                if image is None:
                    image_part = document.part.related_parts.get(rel_id)
                    if image_part is None:
                        logger.warning("[DOCX] 图片关系缺失：%s", rel_id)
                        continue
                    output = _save_image_part(
                        image_part, assets_dir, image_sequence
                    )
                    alt = _alt_text(element, rel_id)
                    try:
                        description = _describe_image(output)
                    except Exception:
                        logger.exception("[DOCX] 图片描述失败：%s", output)
                        description = alt or "[图片描述生成失败]"
                    image = _ImageAsset(str(output.resolve()), description, alt)
                    image_cache[rel_id] = image
                    image_sequence += 1
                block_images.append(image)
            image_refs = {
                rel_id: image_cache[rel_id]
                for rel_id in rel_ids
                if rel_id in image_cache
            }
            if not text.strip() and has_visual_content and not block_images:
                text = "[嵌入式图形或图表]"
            if text.strip() or block_images or has_visual_content:
                rendered_html = _render_paragraph(paragraph, image_refs)
                list_info = _paragraph_list(paragraph)
                if list_info is not None:
                    list_tag = "ol" if list_info[0] == "ordered" else "ul"
                    rendered_html = (
                        f'<{list_tag} data-list-level="{list_info[1]}">'
                        f"<li>{rendered_html}</li></{list_tag}>"
                    )
                page_contents[page_number - 1].append(rendered_html)
        else:
            table_xml = element
            table = Table(table_xml, document)
            row_count = len(table.rows)
            row_groups: list[tuple[int, list[int]]] = []
            for row_index in range(row_count):
                page_number = table_row_pages.get((source_sequence, row_index), seen_page)
                page_number = min(max(page_number, 1), page_count)
                seen_page = max(seen_page, page_number)
                if row_groups and row_groups[-1][0] == page_number:
                    row_groups[-1][1].append(row_index)
                else:
                    row_groups.append((page_number, [row_index]))
            for group_index, (page_number, row_indices) in enumerate(row_groups):
                html, _, image_sequence = _table_html(
                    table, document.part, assets_dir, image_sequence,
                    image_cache, row_indices=row_indices,
                )
                html = html.replace(
                    "<table>",
                    f'<table data-table-id="table_{source_sequence:04d}" '
                    f'data-continues-from-previous="{str(group_index > 0).lower()}" '
                    f'data-continues-on-next="{str(group_index < len(row_groups) - 1).lower()}">',
                    1,
                )
                page_contents[page_number - 1].append(html)

    return tuple(
        DOCXPage(page_number=index + 1, content="".join(content))
        for index, content in enumerate(page_contents)
    )


def parse_docx(file_path: str | Path, *, document_id: str | None = None) -> DOCXDocument:
    """按 DOCX 实际排版分页，返回结构块、表格 HTML 和图片描述。

    图片资源保存到项目根目录 ``data/<document_id>/``。不传 ID 时生成临时目录名；
    解析阶段不计算 MD5。
    """
    try:
        from docx import Document as OpenDocument  # noqa: F401
        import pymupdf  # noqa: F401
    except ImportError as exc:
        raise ImportError("DOCX 页面解析需要 python-docx 和 PyMuPDF；运行 uv sync") from exc

    path = Path(file_path).expanduser()
    if not path.is_file():
        raise FileNotFoundError(f"DOCX 文件不存在：{path}")
    doc_id = document_id or f"pending-{uuid4().hex}"
    if not re.fullmatch(r"[A-Za-z0-9._-]+", doc_id) or doc_id in {".", ".."}:
        raise ValueError("document_id 必须是单个路径片段，例如文档 MD5")

    project_root = Path(__file__).resolve().parent.parent
    assets_dir = project_root / "data" / doc_id
    assets_dir.mkdir(parents=True, exist_ok=True)
    locations = _word_page_boundaries(path)
    if locations is None:
        with tempfile.TemporaryDirectory(prefix="docx-pages-") as temporary:
            pdf_path = _export_docx_pdf(path, Path(temporary))
            paragraph_pages, table_row_pages, page_count = _page_boundaries(path, pdf_path)
    else:
        paragraph_pages, table_row_pages, page_count = locations
    logger.info(
        "[DOCX] 分页定位完成：页数 %d，段落映射 %d，表格行映射 %d",
        page_count,
        len(paragraph_pages),
        len(table_row_pages),
    )
    page_count = max(page_count, 1)
    pages = _document_pages(
        path, assets_dir, paragraph_pages, table_row_pages, page_count
    )

    logger.info(
        "[DOCX解析] 完成：%s，页数 %d，资源目录 data/%s",
        path,
        len(pages),
        doc_id,
    )
    return DOCXDocument(
        source_file=str(path.resolve()),
        pages=pages,
    )


__all__ = ["DOCXDocument", "DOCXPage", "parse_docx"]
